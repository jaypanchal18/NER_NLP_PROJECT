import streamlit as st
import pandas as pd
from concurrent.futures import ThreadPoolExecutor
import io
import nltk
from nltk.tokenize import word_tokenize
from nltk import pos_tag, ne_chunk
from docx import Document
import matplotlib.pyplot as plt
import base64
import openpyxl
import re
import fitz  # PyMuPDF

# Download NLTK resources
nltk.download('punkt')
nltk.download('averaged_perceptron_tagger')
nltk.download('maxent_ne_chunker')
nltk.download('words')

# Handle NaN values in 'Narration' column
def handle_nan(narration):
    return str(narration) if pd.notnull(narration) else ""

# Define extraction functions
def extract_sender_receiver_name(narration):
    names = re.findall(r'\b(?:[A-Z]+\s)+[A-Z]+\b', handle_nan(narration))
    return names[0] if names else None

def extract_payment_method(narration):
    methods = re.findall(r'(?:UPI|IMPS|NEFT|RTGS)\b', handle_nan(narration))
    return methods[0] if methods else None

def extract_payment_platform(narration):
    narration = handle_nan(narration).lower()
    if 'phone pe' in narration or '@ybl' in narration or '@axl' in narration:
        return 'PhonePe'
    elif 'paytm' in narration:
        return 'Paytm'
    elif 'bharatpe' in narration:
        return 'Bharat Pay'
    elif 'cash wdl' in narration:
        return 'ATM'
    elif '@ok' in narration:
        return 'Google Pay'
    else:
        return None

def extract_text_from_pdf(pdf_file):
    text = ""
    pdf_document = fitz.open(pdf_file)
    for page_num in range(len(pdf_document)):
        page = pdf_document.load_page(page_num)
        text += page.get_text()
    return text

def extract_text_from_docx(docx_file):
    doc = Document(docx_file)
    text = ""
    for paragraph in doc.paragraphs:
        text += paragraph.text + "\n"
    return text

def nltk_named_entity_recognition(text):
    tokens = word_tokenize(text)
    tagged = pos_tag(tokens)
    entities = ne_chunk(tagged)
    return entities

def process_excel_file(file):
    wb = openpyxl.load_workbook(file)
    sheet = wb.active
    data = []
    for row in sheet.iter_rows(values_only=True):
        data.append(row)
    df = pd.DataFrame(data[1:], columns=data[0])
    return df

def main():
    st.title('Upload Transaction File')

    uploaded_file = st.file_uploader("Choose a file", type=['csv', 'xlsx', 'pdf', 'docx'])

    if uploaded_file is not None:
        try:
            st.write("Uploaded File Type:", uploaded_file.type)  # Debug statement
            
            if uploaded_file.type == 'text/csv':
                df = pd.read_csv(io.StringIO(uploaded_file.read().decode('utf-8')))
                
            elif uploaded_file.type == 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet':
                df = process_excel_file(uploaded_file)
            
            elif uploaded_file.type == 'application/pdf':
                # Process PDF file
                with ThreadPoolExecutor() as executor:
                    text, entities = executor.submit(process_pdf_file, uploaded_file).result()
                df = pd.DataFrame({'Narration': [text], 'Named_Entities': [entities]})
            elif uploaded_file.type == 'application/vnd.openxmlformats-officedocument.wordprocessingml.document':
                # Process DOCX file
                with ThreadPoolExecutor() as executor:
                    text, entities = executor.submit(process_docx_file, uploaded_file).result()
                df = pd.DataFrame({'Narration': [text], 'Named_Entities': [entities]})
            else:
                st.error('Unsupported file format')

            # Data processing
            if 'Amount' in df.columns:
                total_amount = df['Amount'].sum()
            else:
                total_amount = 'N/A (Amount column not found)'

            df['Name'] = df['Narration'].apply(extract_sender_receiver_name)
            df['Payment_Method'] = df['Narration'].apply(extract_payment_method)
            df['Payment_Platform'] = df['Narration'].apply(extract_payment_platform)


            # Visualization
            st.subheader('Extracted Information')
            st.write(df)

            st.subheader(f"**Total Amount:** {total_amount}")

            # Payment Method Pie Chart
            payment_method_counts = df['Payment_Method'].value_counts()
            fig_pie, ax_pie = plt.subplots(figsize=(5, 4))
            ax_pie.pie(payment_method_counts, labels=payment_method_counts.index, autopct='%1.1f%%', startangle=140)
            ax_pie.set_title('Payment Method Distribution')
            st.pyplot(fig_pie)

            # Payment Platform Bar Chart
            payment_platform_counts = df['Payment_Platform'].value_counts()
            fig_bar, ax_bar = plt.subplots(figsize=(6, 4))
            payment_platform_counts.plot(kind='bar', color='skyblue', ax=ax_bar)
            ax_bar.set_title('Payment Platform Distribution')
            ax_bar.set_xlabel('Payment Platform')
            ax_bar.set_ylabel('Frequency')
            ax_bar.tick_params(axis='x', rotation=45)
            st.pyplot(fig_bar)

            # Cash Withdrawal Pie Chart
            cash_withdrawals = (df['Payment_Platform'] == 'ATM').sum()
            other_transactions = len(df) - cash_withdrawals
            fig_cash_withdrawal, ax_cash_withdrawal = plt.subplots()
            ax_cash_withdrawal.pie([cash_withdrawals, other_transactions], labels=['Cash Withdrawal', 'Other Transactions'], autopct='%1.1f%%')
            ax_cash_withdrawal.set_title('Cash Withdrawal vs Other Transactions')
            st.pyplot(fig_cash_withdrawal)

            # Download buttons
            st.subheader('Download Data')
            data = df.to_csv(index=False)
            csv_button = st.download_button(label='Download CSV', data=data, file_name='data.csv', mime='text/csv')

            pdf_button = st.download_button(label='Download PDF', data=data, file_name='data.pdf', mime='application/pdf')

            xlsx_button = st.download_button(label='Download Excel', data=data, file_name='data.xlsx', mime='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
            

        except Exception as e:
            st.error(f'Error processing file: {str(e)}')

if __name__ == '__main__':
    main()
